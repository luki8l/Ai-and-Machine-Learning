"""Builds LSEG_IV_Download.xlsx: one sheet per Dow stock with a ready RDP.Data formula."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import sys
sys.path.insert(0, "..")
from volml import DOW30_VOLARE

CALL = "TR.30DAYATTHEMONEYIMPLIEDVOLATILITYINDEXFORCALLOPTIONS"
PUT = "TR.30DAYATTHEMONEYIMPLIEDVOLATILITYINDEXFORPUTOPTIONS"
START, END = "2016-01-01", "2026-10-16"

wb = Workbook()
ws = wb.active
ws.title = "Anleitung"
lines = [
    ("LSEG Implied Volatility Download (30-Tage ATM, Call + Put), Dow 30, 2016 bis heute", True),
    ("", False),
    ("1. LSEG Workspace öffnen und einloggen (Uni-Account).", False),
    ("2. Diese Datei in Excel öffnen. Der Reiter 'LSEG' bzw. 'Workspace' muss sichtbar sein, dort ggf. einloggen.", False),
    ("3. Im LSEG-Reiter auf 'Refresh' / 'Refresh All' klicken (oder Strg+Alt+F9).", False),
    ("   Zeigt eine Zelle #NAME?: Zelle A3 auf dem Blatt anklicken, F2 drücken, Enter. Dann lädt die Formel.", False),
    ("4. Warten, bis alle 30 Blätter gefüllt sind. Blatt 'Kontrolle' zeigt pro Aktie die Anzahl Zeilen (Ziel: ca. 2.700).", False),
    ("5. ALLES ALS WERTE SPEICHERN: Datei > Speichern unter > neue Datei 'LSEG_IV_values.xlsx'.", False),
    ("   Dann in der neuen Datei auf jedem Aktien-Blatt: Strg+A, Strg+C, Rechtsklick > Einfügen als Werte.", False),
    ("   (Alternative: LSEG-Reiter > 'Convert to values' / 'Remove formulas', falls vorhanden.)", False),
    ("6. Im Projektordner: python lseg/convert_lseg_excel.py LSEG_IV_values.xlsx  -> erzeugt data/lseg/iv.csv", False),
    ("   Oder die Datei an Claude schicken.", False),
    ("", False),
    ("Hinweise", True),
    ("- RIC = Ticker + 'ATMIV.U' (z.B. AAPLATMIV.U). Felder laut LSEG Developer Forum, hier nicht getestet.", False),
    ("- Zeigt ein Blatt einen Fehler/#N/A: RIC in Workspace suchen. Fehlende Aktien sind ok, IV wird dann nur für die übrigen genutzt.", False),
    ("- Bricht eine Abfrage ab (Limit), auf dem Blatt in der Formel SDate z.B. auf 2021-01-01 setzen und den Rest auf einem Kopie-Blatt holen.", False),
    (f"- Formel-Aufbau: =RDP.Data(RIC; \"{CALL}.date;{CALL};{PUT}\"; \"Frq=D SDate={START} EDate={END} CH=Fd\")", False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=12 if bold else 11)
ws.column_dimensions["A"].width = 140

ctrl = wb.create_sheet("Kontrolle")
ctrl.append(["Ticker", "RIC", "Zeilen mit Call-IV", "Zeilen mit Put-IV"])
for c in ctrl[1]:
    c.font = Font(bold=True)

for t in DOW30_VOLARE:
    s = wb.create_sheet(t)
    ric = f"{t}ATMIV.U"
    s["A1"] = f"{t}  |  {ric}"
    s["A1"].font = Font(bold=True)
    s["A1"].fill = PatternFill("solid", fgColor="DDEBF7")
    s["A3"] = (f'=RDP.Data("{ric}","{CALL}.date;{CALL};{PUT}",'
               f'"Frq=D SDate={START} EDate={END} CH=Fd")')
    for col, w in (("A", 14), ("B", 14), ("C", 14)):
        s.column_dimensions[col].width = w
    r = ctrl.max_row + 1
    ctrl.cell(row=r, column=1, value=t)
    ctrl.cell(row=r, column=2, value=ric)
    ctrl.cell(row=r, column=3, value=f"=COUNT('{t}'!B:B)")
    ctrl.cell(row=r, column=4, value=f"=COUNT('{t}'!C:C)")
for col, w in (("A", 10), ("B", 16), ("C", 20), ("D", 20)):
    ctrl.column_dimensions[col].width = w
wb.save("LSEG_IV_Download.xlsx")
print("ok")
